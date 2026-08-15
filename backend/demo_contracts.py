"""Contrats DOCX synthétiques pour une démonstration sans donnée sensible."""

from io import BytesIO
from typing import Literal

from docx import Document


DemoVersion = Literal["before", "after"]


DEMO_CONTENT: dict[DemoVersion, list[tuple[str, str]]] = {
    "before": [
        (
            "Confidentialité",
            "Chaque partie respecte une obligation de confidentialité pendant deux ans et s'interdit toute divulgation à un tiers.",
        ),
        (
            "Non-concurrence",
            "Une clause de non-concurrence interdit au prestataire toute activité concurrente pendant douze mois en France.",
        ),
    ],
    "after": [
        (
            "Confidentialité",
            "Chaque partie respecte une obligation de confidentialité pendant cinq ans et s'interdit toute divulgation à un tiers.",
        ),
        (
            "Restrictions commerciales",
            "Les parties confirment qu'aucune clause de non-concurrence n'est prévue dans cette version.",
        ),
        (
            "Gouvernance",
            "Le conseil d'administration statue à la majorité qualifiée, sous réserve d'un quorum de soixante pour cent.",
        ),
    ],
}


def build_demo_contract(version: DemoVersion) -> bytes:
    """Génère en mémoire un document qui ne contient aucune donnée réelle."""

    document = Document()
    document.add_heading(
        f"Contrat de démonstration — {'version précédente' if version == 'before' else 'nouvelle version'}",
        level=0,
    )
    document.add_paragraph(
        "Document entièrement synthétique généré par ClauseScope. Ne contient aucune donnée contractuelle réelle."
    )
    for title, text in DEMO_CONTENT[version]:
        document.add_heading(title, level=1)
        document.add_paragraph(text)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
